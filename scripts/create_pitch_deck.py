import os
import sys
import qrcode
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Define color palette
COLOR_BG = RGBColor(3, 7, 18)           # #030712 Near Black
COLOR_CARD = RGBColor(11, 15, 25)        # #0B0F19 Deep Charcoal
COLOR_CARD_BORDER = RGBColor(30, 41, 59) # #1E293B Subtle Border
COLOR_CYAN = RGBColor(0, 240, 255)       # #00F0FF Electric Cyan
COLOR_GREEN = RGBColor(34, 197, 94)      # #22C55E Emerald Verified
COLOR_AMBER = RGBColor(245, 158, 11)     # #F59E0B Amber Warning
COLOR_WHITE = RGBColor(255, 255, 255)    # #FFFFFF Crisp White
COLOR_MUTED = RGBColor(148, 163, 184)    # #94A3B8 Slate Gray
COLOR_RED_MUTED = RGBColor(239, 68, 68)  # #EF4444 Muted Red

FONT_HEADING = "Arial"
FONT_BODY = "Arial"
FONT_MONO = "Consolas"

SCREENSHOT_DIR = "C:/Users/heman/.gemini/antigravity-ide/brain/96d0904b-ffea-4253-98c1-7ef114a5795d"

def set_slide_background(slide, prs):
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG
    bg.line.fill.background()
    return bg

def add_header(slide, tag_text, title_text, subtitle_text=None):
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tag_text.upper()
    p_tag.font.name = FONT_MONO
    p_tag.font.size = Pt(10)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_CYAN

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.75))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(32)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE

    if subtitle_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.4))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle_text
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(14)
        p_sub.font.color.rgb = COLOR_MUTED

def add_card(slide, left, top, width, height, fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = fill_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card

def add_speaker_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text

def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Generate QR Code for https://verifyx-gray.vercel.app/
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=10,
        border=2,
    )
    qr.add_data("https://verifyx-gray.vercel.app/")
    qr.make(fit=True)
    qr_img = qr.make_image(fill_color="black", back_color="white")
    qr_path = os.path.join(SCREENSHOT_DIR, "verifyx_live_qr.png")
    qr_img.save(qr_path)

    # ====================================================
    # SLIDE 1 — HERO
    # ====================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, prs)

    tag1 = s1.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(6.0), Inches(0.35))
    p = tag1.text_frame.paragraphs[0]
    p.text = "STAGE 01 // DEPLOYED APPLICATION · LIVE VERIFICATION ENGINE"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    pill = add_card(s1, Inches(0.8), Inches(1.1), Inches(2.2), Inches(0.35), fill_color=RGBColor(15, 23, 42), border_color=COLOR_CYAN)
    p_pill = pill.text_frame.paragraphs[0]
    p_pill.text = "● LIVE ON VERCEL"
    p_pill.alignment = PP_ALIGN.CENTER
    p_pill.font.name = FONT_MONO
    p_pill.font.size = Pt(10)
    p_pill.font.bold = True
    p_pill.font.color.rgb = COLOR_CYAN

    t1 = s1.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(5.8), Inches(1.2))
    p = t1.text_frame.paragraphs[0]
    p.text = "VERIFYX"
    p.font.name = FONT_HEADING
    p.font.size = Pt(64)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE

    sub1 = s1.shapes.add_textbox(Inches(0.8), Inches(2.85), Inches(5.8), Inches(0.6))
    p = sub1.text_frame.paragraphs[0]
    p.text = '"Verify Before You Trust."'
    p.font.name = FONT_HEADING
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    sup1 = s1.shapes.add_textbox(Inches(0.8), Inches(3.55), Inches(5.6), Inches(1.5))
    tf = sup1.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "AI-powered verification for detecting suspicious or potentially non-compliant physical workplace conditions."
    p1.font.name = FONT_BODY
    p1.font.size = Pt(14)
    p1.font.color.rgb = COLOR_WHITE
    
    p2 = tf.add_paragraph()
    p2.space_before = Pt(12)
    p2.text = "A phone-first, in-browser AI verification engine. Reduces visual inspection ambiguity through automated object detection, deterministic spatial rule evaluation, and closed-loop audit reports."
    p2.font.name = FONT_BODY
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_MUTED

    m_box = s1.shapes.add_textbox(Inches(0.8), Inches(5.5), Inches(5.8), Inches(0.5))
    p = m_box.text_frame.paragraphs[0]
    p.text = "24/24 UNIT TESTS PASS  •  DEVICE-LOCAL PROCESSING  •  ASYNC ~2 FPS HUD"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN

    f1 = s1.shapes.add_textbox(Inches(0.8), Inches(6.6), Inches(11.7), Inches(0.4))
    p = f1.text_frame.paragraphs[0]
    p.text = "AI • VERIFICATION • WORKFLOW    |    DEPLOYED: https://verifyx-gray.vercel.app/"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_MUTED

    img_home = os.path.join(SCREENSHOT_DIR, "home_screen_1790082836407.png")
    if os.path.exists(img_home):
        add_card(s1, Inches(6.8), Inches(1.1), Inches(5.7), Inches(5.0), fill_color=COLOR_CARD, border_color=COLOR_CYAN)
        s1.shapes.add_picture(img_home, Inches(6.9), Inches(1.2), Inches(5.5), Inches(4.8))

    add_speaker_notes(s1, """Judges, good morning. In an era where visual documentation and physical compliance are easily misdiagnosed or overlooked, trusting raw visual inspection alone creates massive operational blind spots.

This is VERIFYX — built around a straightforward operating principle: Verify Before You Trust.

VERIFYX is not a slide concept, not a mockup, and not an unvalidated research paper. It is an active, deployed web application running right now on Vercel at verifyx-gray.vercel.app. It performs visual object detection, deterministic spatial reasoning, and closed-loop verification directly in the user's browser. Let me show you how it works.""")

    # ====================================================
    # SLIDE 2 — THE PROBLEM
    # ====================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, prs)
    add_header(s2, "STAGE 02 // MACRO CONTEXT · THE TRUST CRISIS", '"Seeing is no longer believing."', "The modern verification challenge across visual environments")

    flow_box = add_card(s2, Inches(0.8), Inches(1.9), Inches(11.7), Inches(0.8), fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER)
    tf = flow_box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "CONTENT   ───>   UNKNOWN ORIGIN   ───>   MANIPULATION / INSPECTION RISK   ───>   UNCERTAIN TRUST"
    p.font.name = FONT_MONO
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_AMBER

    card_w = Inches(3.7)
    card_h = Inches(2.9)
    top_pos = Inches(2.9)

    c1 = add_card(s2, Inches(0.8), top_pos, card_w, card_h)
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "01\nINFORMATION OVERLOAD"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p_body = tf1.add_paragraph()
    p_body.space_before = Pt(14)
    p_body.text = "Facility managers, safety operators, and auditors encounter high volumes of visual scenarios and documentation daily.\n\nThe sheer volume and speed of modern operations can easily outpace manual review capacity."
    p_body.font.name = FONT_BODY
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = COLOR_MUTED

    c2 = add_card(s2, Inches(4.8), top_pos, card_w, card_h)
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "02\nPERCEPTION LIMIT"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN
    p_body = tf2.add_paragraph()
    p_body.space_before = Pt(14)
    p_body.text = "Human visual inspection can struggle with precise spatial relationships, clearance margins, and subtle hazard proximity.\n\nSubjective human review suffers from fatigue, confirmation bias, and inconsistent evaluations."
    p_body.font.name = FONT_BODY
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = COLOR_MUTED

    c3 = add_card(s2, Inches(8.8), top_pos, card_w, card_h)
    tf3 = c3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "03\nFRAGMENTED VERIFICATION"
    p.font.name = FONT_HEADING
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p_body = tf3.add_paragraph()
    p_body.space_before = Pt(14)
    p_body.text = "Photos, notes, physical checklists, and audit reports frequently live in separate, disconnected workflows.\n\nWithout a continuous chain of evidence, verification details are lost between field inspection and sign-off."
    p_body.font.name = FONT_BODY
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = COLOR_MUTED

    b_stmt = add_card(s2, Inches(0.8), Inches(6.05), Inches(11.7), Inches(0.8), fill_color=COLOR_CARD, border_color=COLOR_CYAN)
    tf_b = b_stmt.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.alignment = PP_ALIGN.CENTER
    p_b.text = '"The problem isn\'t finding content. It\'s knowing whether to trust it."'
    p_b.font.name = FONT_HEADING
    p_b.font.size = Pt(16)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_WHITE

    add_speaker_notes(s2, """Every day, operations teams evaluate critical visual evidence: Is that exit corridor truly clear of obstructions? Is emergency equipment present and properly mounted? Has a previously flagged hazard actually been cleared?

Today, visual inspection fails because seeing is no longer believing. Humans are overloaded with visual checks, unable to reliably estimate spatial overlap percentages with the naked eye, and forced to manually copy-paste photos into disconnected spreadsheets.

When verification is fragmented, compliance degrades into guesswork.""")

    # ====================================================
    # SLIDE 3 — THE INSIGHT
    # ====================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, prs)
    add_header(s3, "STAGE 03 // ARCHITECTURAL PARADIGM SHIFT", '"Verification should be a workflow, not a guess."', "Transforming visual verification from subjective inspection to structured evidence")

    c_before = add_card(s3, Inches(0.8), Inches(2.0), Inches(5.6), Inches(3.6), fill_color=RGBColor(15, 12, 18), border_color=COLOR_RED_MUTED)
    tf_bef = c_before.text_frame
    tf_bef.word_wrap = True
    p = tf_bef.paragraphs[0]
    p.text = "BEFORE  //  THE MANUAL GUESS"
    p.font.name = FONT_MONO
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED_MUTED

    p2 = tf_bef.add_paragraph()
    p2.space_before = Pt(20)
    p2.text = "Capture  ──>  Inspect  ──>  Guess"
    p2.font.name = FONT_MONO
    p2.font.size = Pt(16)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_WHITE

    p3 = tf_bef.add_paragraph()
    p3.space_before = Pt(20)
    p3.text = "• Uncalibrated visual appraisal\n• Black-box subjective judgement\n• No spatial dimension calculations\n• Verbal remediation instructions left unverified\n• High variance between individual inspectors"
    p3.font.name = FONT_BODY
    p3.font.size = Pt(13)
    p3.font.color.rgb = COLOR_MUTED

    c_after = add_card(s3, Inches(6.9), Inches(2.0), Inches(5.6), Inches(3.6), fill_color=RGBColor(8, 22, 28), border_color=COLOR_CYAN)
    tf_aft = c_after.text_frame
    tf_aft.word_wrap = True
    p = tf_aft.paragraphs[0]
    p.text = "AFTER  //  THE VERIFYX WORKFLOW"
    p.font.name = FONT_MONO
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    p2 = tf_aft.add_paragraph()
    p2.space_before = Pt(20)
    p2.text = "CAPTURE ──> ANALYZE ──> VERIFY ──> EXPLAIN ──> REPORT"
    p2.font.name = FONT_MONO
    p2.font.size = Pt(12)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_GREEN

    p3 = tf_aft.add_paragraph()
    p3.space_before = Pt(20)
    p3.text = "• Automated local object detection (YOLO11n ONNX)\n• Geometric spatial rule evaluation (IoU overlap)\n• Explicit explainability formula for every verdict\n• Enforced Before/After rescan evidence for remediation\n• Structured, timestamped JSON inspection records"
    p3.font.name = FONT_BODY
    p3.font.size = Pt(13)
    p3.font.color.rgb = COLOR_WHITE

    b_banner = add_card(s3, Inches(0.8), Inches(5.85), Inches(11.7), Inches(1.0), fill_color=COLOR_CARD, border_color=COLOR_GREEN)
    tf_ban = b_banner.text_frame
    p_b1 = tf_ban.paragraphs[0]
    p_b1.text = "PERCEPTION + REASONING = ACTIONABLE DECISION"
    p_b1.font.name = FONT_MONO
    p_b1.font.size = Pt(13)
    p_b1.font.bold = True
    p_b1.font.color.rgb = COLOR_CYAN

    p_b2 = tf_ban.add_paragraph()
    p_b2.space_before = Pt(4)
    p_b2.text = '"VERIFYX turns raw visual input into a structured, explainable decision."'
    p_b2.font.name = FONT_HEADING
    p_b2.font.size = Pt(16)
    p_b2.font.bold = True
    p_b2.font.color.rgb = COLOR_WHITE

    add_speaker_notes(s3, """Our core insight is simple: Verification cannot be a one-off guess. It must be a continuous, structured workflow.

In legacy workflows, someone takes a photo, uploads it, squints at the image, and makes an uncalibrated call.

VERIFYX transforms this into a 5-stage closed-loop system: Capture the scene, Analyze bounding boxes and class detections, Verify against explicit geometric spatial rules, Explain the underlying decision formula, and export a structured Report.

We don't just output a probability score; we provide an explainable verdict.""")

    # ====================================================
    # SLIDE 4 — THE SOLUTION
    # ====================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, prs)
    add_header(s4, "STAGE 04 // PRODUCT CAPABILITIES · LIVE SYSTEM", "Meet VERIFYX.", "An AI-powered verification workflow that turns ambiguous visual inputs into an explainable verification result.")

    caps = [
        ("01 LIVE SCAN", "Real-time camera feed with in-browser inference reticles throttled at ~2 FPS async."),
        ("02 DATASET DEMO", "Pre-loaded evaluation suite running 5 standardized workplace safety test cases."),
        ("03 VERIFICATION RESULTS", "Real-time checklist resolution tallying verified standards (4 / 5 VERIFIED)."),
        ("04 ISSUE DETECTION", "Mathematical spatial rule engine identifying bounding box conflicts & pathway intrusions."),
        ("05 RESCAN", "Closed-loop remediation workflow: side-by-side Before/After targeting confirming clearance."),
        ("06 AUDIT REPORT", "Structured inspection log (VX-INSP-2026-XXXX) with timestamps and JSON export.")
    ]

    for i, (title, desc) in enumerate(caps):
        row = i // 2
        col = i % 2
        c_left = Inches(0.8 + col * 3.3)
        c_top = Inches(2.0 + row * 1.35)
        card = add_card(s4, c_left, c_top, Inches(3.15), Inches(1.2), fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER)
        tf_c = card.text_frame
        tf_c.word_wrap = True
        p1 = tf_c.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_MONO
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_CYAN
        p2 = tf_c.add_paragraph()
        p2.space_before = Pt(3)
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(10)
        p2.font.color.rgb = COLOR_MUTED

    pipe_box = add_card(s4, Inches(0.8), Inches(6.25), Inches(6.45), Inches(0.65), fill_color=COLOR_CARD, border_color=COLOR_CYAN)
    tf_p = pipe_box.text_frame
    p = tf_p.paragraphs[0]
    p.text = "INPUT  ──>  AI ANALYSIS  ──>  SPATIAL RULES  ──>  VERIFICATION RESULT  ──>  REPORT"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_MONO
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE

    img_demo = os.path.join(SCREENSHOT_DIR, "dataset_demo_fe01_1790083347828.png")
    img_results = os.path.join(SCREENSHOT_DIR, "verification_results_4of5_1790084431011.png")

    if os.path.exists(img_demo):
        add_card(s4, Inches(7.5), Inches(2.0), Inches(5.0), Inches(2.35), fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER)
        s4.shapes.add_picture(img_demo, Inches(7.55), Inches(2.05), Inches(4.9), Inches(2.25))

    if os.path.exists(img_results):
        add_card(s4, Inches(7.5), Inches(4.55), Inches(5.0), Inches(2.35), fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER)
        s4.shapes.add_picture(img_results, Inches(7.55), Inches(4.6), Inches(4.9), Inches(2.25))

    add_speaker_notes(s4, """Meet VERIFYX.

This is a functional, end-to-end verification engine. As you can see, VERIFYX allows an operator to point their device camera at a target scene — or run through our built-in dataset evaluation mode.

The engine scans the frame, isolates bounding boxes, checks spatial overlaps against safety rules, flags specific issues, requires a rescan to prove remediation, and packages the results into an inspection report.

You can try this right now at verifyx-gray.vercel.app.""")

    # ====================================================
    # SLIDE 5 — HOW IT WORKS
    # ====================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, prs)
    add_header(s5, "STAGE 05 // SYSTEM ARCHITECTURE & PIPELINE", "Inside the VERIFYX verification pipeline.", "A 4-tier decoupled architecture combining in-browser neural perception and deterministic spatial logic")

    layers = [
        ("USER LAYER", "WebRTC Camera Stream (640x640 Canvas)  •  Dataset Evaluation Suite (5 Test Scenarios)", COLOR_WHITE),
        ("PROCESSING LAYER", "Frame Normalization (RGB, Float32)  •  Temporal Confirmation (Sliding Window Buffer)", COLOR_MUTED),
        ("INTELLIGENCE LAYER", "Primary: Custom YOLO11n ONNX via WebAssembly (onnxruntime-web 1.30.0)\nFallback: Local TF.js COCO-SSD  ·  Deterministic Simulation Engine\nSpatial Engine: Bounding Box Overlap, Size Guardrail (<32px), Egress Pathway Clearance", COLOR_CYAN),
        ("OUTPUT LAYER", "Dynamic Reticle HUD  •  Issue Explainability Matrix  •  Structured JSON Report (VX-INSP-2026-XXXX)", COLOR_GREEN)
    ]

    for i, (title, details, accent) in enumerate(layers):
        l_top = Inches(2.0 + i * 1.08)
        card = add_card(s5, Inches(0.8), l_top, Inches(11.7), Inches(0.95), fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER)
        
        pill_box = s5.shapes.add_textbox(Inches(1.0), l_top + Inches(0.08), Inches(4.5), Inches(0.3))
        tf_pill = pill_box.text_frame
        tf_pill.margin_left = tf_pill.margin_top = tf_pill.margin_right = tf_pill.margin_bottom = 0
        p1 = tf_pill.paragraphs[0]
        p1.text = title.upper()
        p1.font.name = FONT_MONO
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = accent

        det_box = s5.shapes.add_textbox(Inches(1.0), l_top + Inches(0.38), Inches(11.3), Inches(0.5))
        tf_det = det_box.text_frame
        tf_det.word_wrap = True
        tf_det.margin_left = tf_det.margin_top = tf_det.margin_right = tf_det.margin_bottom = 0
        p2 = tf_det.paragraphs[0]
        p2.text = details
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_WHITE

    priv_box = add_card(s5, Inches(0.8), Inches(6.45), Inches(11.7), Inches(0.6), fill_color=COLOR_CARD, border_color=COLOR_CYAN)
    tf_pr = priv_box.text_frame
    p = tf_pr.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "TECHNICAL PRINCIPLE: Camera frames are processed locally in the browser; they are not transmitted to external servers."
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    add_speaker_notes(s5, """Let's look at the technical architecture that powers VERIFYX.

The system operates across four decoupled layers:
First, the User Layer acquires video via WebRTC canvas streams.
Second, the Processing Layer normalizes pixel tensors and applies a lightweight temporal confirmation buffer across consecutive frames.
Third, the Intelligence Layer: our primary engine runs a custom YOLO11n ONNX model executing inside browser WebAssembly via ONNX Runtime Web. It detects target objects, while our pure TypeScript Spatial Rule Engine evaluates bounding box overlap, minimum pixel dimensions, and clearance zones. If WASM encounters restrictions, it gracefully falls back to local TensorFlow.js or our deterministic simulation adapter.
Finally, the Output Layer renders interactive HUD overlays and generates structured JSON reports.

All image processing happens locally on the client device without sending video frames to cloud servers.""")

    # ====================================================
    # SLIDE 6 — PRODUCT EXPERIENCE
    # ====================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, prs)
    add_header(s6, "STAGE 06 // USER WORKFLOW & INTERACTION DESIGN", "From scan to decision in one flow.", "The complete closed-loop inspection journey: point, verify, fix, and verify again")

    steps = [
        ("01 LIVE SCAN", "Real-time camera stream HUD"),
        ("02 ANALYZE", "In-browser YOLO inference"),
        ("03 RESULT", "Instant 4/5 checklist resolution"),
        ("04 ISSUE DETAILS", "Explicit decision formula"),
        ("05 RESCAN", "Before/After clearance check"),
        ("06 REPORT", "Structured audit log & export")
    ]

    for i, (st_num, st_desc) in enumerate(steps):
        s_left = Inches(0.8 + i * 1.98)
        c_step = add_card(s6, s_left, Inches(1.9), Inches(1.85), Inches(0.9), fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER)
        tf_s = c_step.text_frame
        tf_s.word_wrap = True
        p1 = tf_s.paragraphs[0]
        p1.text = st_num
        p1.font.name = FONT_MONO
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_CYAN
        p2 = tf_s.add_paragraph()
        p2.space_before = Pt(2)
        p2.text = st_desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(9)
        p2.font.color.rgb = COLOR_MUTED

    form_box = add_card(s6, Inches(0.8), Inches(2.95), Inches(11.7), Inches(0.65), fill_color=COLOR_CARD, border_color=COLOR_AMBER)
    tf_f = form_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.alignment = PP_ALIGN.CENTER
    p_f.text = "EXPLAINABLE FORMULA:  DETECTION [Chair @ (x,y,w), Conf: 88%]  +  SPATIAL RULE [Pathway Overlap > 10%]  =  POSSIBLE OBSTRUCTION"
    p_f.font.name = FONT_MONO
    p_f.font.size = Pt(10)
    p_f.font.bold = True
    p_f.font.color.rgb = COLOR_WHITE

    shots = [
        ("live_scan_screen_1790083147852.png", "Live Targeting"),
        ("issue_details_screen_1790084564301.png", "Reasoning Formula"),
        ("rescan_screen_1790084662137.png", "Closed-Loop Rescan"),
        ("report_screen_1790084943630.png", "Audit Record")
    ]

    for i, (img_name, label) in enumerate(shots):
        i_left = Inches(0.8 + i * 2.98)
        img_path = os.path.join(SCREENSHOT_DIR, img_name)
        if os.path.exists(img_path):
            add_card(s6, i_left, Inches(3.75), Inches(2.8), Inches(3.3), fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER)
            s6.shapes.add_picture(img_path, i_left + Inches(0.05), Inches(3.8), Inches(2.7), Inches(2.9))
            lbl_box = s6.shapes.add_textbox(i_left, Inches(6.75), Inches(2.8), Inches(0.3))
            p_lbl = lbl_box.text_frame.paragraphs[0]
            p_lbl.alignment = PP_ALIGN.CENTER
            p_lbl.text = label.upper()
            p_lbl.font.name = FONT_MONO
            p_lbl.font.size = Pt(9)
            p_lbl.font.bold = True
            p_lbl.font.color.rgb = COLOR_CYAN

    add_speaker_notes(s6, """An AI model alone is not a solution. A complete solution is a user workflow that resolves an inspection task without friction.

In VERIFYX, an operator completes a verification cycle in a single coherent flow:
They run a Live Scan, the engine evaluates the scene, displays a 4/5 checklist result, breaks down why item #4 was flagged, prompts an immediate physical correction and rescan, and generates an inspection record.

We closed the loop between detecting an issue and verifying that it was corrected.""")

    # ====================================================
    # SLIDE 7 — THE DIFFERENTIATOR
    # ====================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, prs)
    add_header(s7, "STAGE 07 // WORKFLOW ADVANTAGE & POSITIONING", '"VERIFYX doesn\'t just flag. It explains."', "Objective architectural comparison between manual inspection methods and VERIFYX")

    hdr_box = add_card(s7, Inches(0.8), Inches(2.0), Inches(11.7), Inches(0.55), fill_color=COLOR_CARD, border_color=COLOR_CYAN)
    
    col1_hdr = s7.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(2.2), Inches(0.35))
    p = col1_hdr.text_frame.paragraphs[0]
    p.text = "DIMENSION"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    col2_hdr = s7.shapes.add_textbox(Inches(3.3), Inches(2.1), Inches(4.3), Inches(0.35))
    p = col2_hdr.text_frame.paragraphs[0]
    p.text = "MANUAL INSPECTION METHODS"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_MUTED

    col3_hdr = s7.shapes.add_textbox(Inches(7.7), Inches(2.1), Inches(4.6), Inches(0.35))
    p = col3_hdr.text_frame.paragraphs[0]
    p.text = "THE VERIFYX WORKFLOW"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN

    rows = [
        ("Tooling", "Disconnected: Camera + Paper / Spreadsheets", "Single In-Browser Web Application"),
        ("Analysis", "Subjective, uncalibrated visual estimate", "Object Detection + Geometric Spatial Rules"),
        ("Explainability", "Intuition-based ('looks blocked')", "Explicit Breakdown: Detection + Spatial Rule = Verdict"),
        ("Remediation", "Untracked verbal instruction", "Enforced Before/After Rescan Evidence"),
        ("Audit Trail", "Fragmented photos & unlinked notes", "Structured JSON with Session ID & Timestamps"),
        ("Processing", "Manual review after the inspection", "Real-Time On-Device Assessment")
    ]

    for i, (dim, manual, verifyx) in enumerate(rows):
        r_top = Inches(2.65 + i * 0.58)
        row_card = add_card(s7, Inches(0.8), r_top, Inches(11.7), Inches(0.5), fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER)
        
        c1_box = s7.shapes.add_textbox(Inches(1.0), r_top + Inches(0.08), Inches(2.2), Inches(0.35))
        p1 = c1_box.text_frame.paragraphs[0]
        p1.text = dim
        p1.font.name = FONT_BODY
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_WHITE

        c2_box = s7.shapes.add_textbox(Inches(3.3), r_top + Inches(0.08), Inches(4.3), Inches(0.35))
        p2 = c2_box.text_frame.paragraphs[0]
        p2.text = manual
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_MUTED

        c3_box = s7.shapes.add_textbox(Inches(7.7), r_top + Inches(0.08), Inches(4.6), Inches(0.35))
        p3 = c3_box.text_frame.paragraphs[0]
        p3.text = verifyx
        p3.font.name = FONT_BODY
        p3.font.size = Pt(11)
        p3.font.bold = True
        p3.font.color.rgb = COLOR_CYAN

    bot_box = add_card(s7, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.7), fill_color=COLOR_CARD, border_color=COLOR_GREEN)
    tf_bot = bot_box.text_frame
    p = tf_bot.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "CORE ADVANTAGE:  DETECTION + SPATIAL RULE = EXPLAINED VERDICT"
    p.font.name = FONT_MONO
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN

    add_speaker_notes(s7, """Standard computer vision demos often present raw bounding boxes and confidence percentages. To an inspection or safety team, a raw probability number is not an actionable decision.

VERIFYX is built differently: We don't just flag; we explain.
When an obstruction is detected, VERIFYX presents the underlying formula: Detected Object, plus Pathway Zone overlap, plus Rule Threshold equals Required Action.

It transforms subjective visual assessments into structured, explainable findings.""")

    # ====================================================
    # SLIDE 8 — ENGINEERING & VALIDATION
    # ====================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, prs)
    add_header(s8, "STAGE 08 // TECHNICAL CREDIBILITY & TEST RIGOR", "Built. Tested. Deployed.", "Demonstrated software engineering discipline with zero mockups and 100% test pass rate")

    hero_card = add_card(s8, Inches(0.8), Inches(2.0), Inches(3.8), Inches(4.8), fill_color=COLOR_CARD, border_color=COLOR_GREEN)
    tf_hero = hero_card.text_frame
    tf_hero.word_wrap = True
    p1 = tf_hero.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    p1.text = "24 / 24"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(54)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_GREEN

    p2 = tf_hero.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "TESTS PASSED (100%)"
    p2.font.name = FONT_MONO
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_WHITE

    p3 = tf_hero.add_paragraph()
    p3.space_before = Pt(16)
    p3.text = "• 13 Spatial Rule Engine Tests\n• 11 Report Lifecycle Tests\n• Edge case validation (<32px guardrail)\n• Temporal buffer tests\n• No fabricated evidence test check\n• Dev mock transfer bridge test"
    p3.font.name = FONT_BODY
    p3.font.size = Pt(11)
    p3.font.color.rgb = COLOR_MUTED

    val_items = [
        ("BUILD STATUS", "PASS (Zero Errors)"),
        ("MOBILE VIEWPORT", "VALIDATED\nTouch HUD"),
        ("DESKTOP VIEWPORT", "VALIDATED\nResponsive"),
        ("LIVE SCAN", "PASS (~2 FPS Async)"),
        ("DATASET DEMO", "PASS (5 Test Cases)"),
        ("CLOSED RESCAN", "PASS (Remediation)"),
        ("AUDIT REPORT", "PASS (JSON Export)"),
        ("GIT REPOSITORY", "Commit afb0952"),
        ("CLOUD HOSTING", "Vercel Edge"),
        ("LIVE APPLICATION", "verifyx-gray.vercel.app")
    ]

    for i, (k, v) in enumerate(val_items):
        row = i // 2
        col = i % 2
        v_left = Inches(4.9 + col * 2.3)
        v_top = Inches(2.0 + row * 0.95)
        v_card = add_card(s8, v_left, v_top, Inches(2.18), Inches(0.85), fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER)
        tf_v = v_card.text_frame
        tf_v.word_wrap = True
        p1 = tf_v.paragraphs[0]
        p1.text = k
        p1.font.name = FONT_MONO
        p1.font.size = Pt(9)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_CYAN
        p2 = tf_v.add_paragraph()
        p2.space_before = Pt(2)
        p2.text = v
        p2.font.name = FONT_HEADING
        p2.font.size = Pt(9.5)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_WHITE

    img_v5 = os.path.join(SCREENSHOT_DIR, "verified_screen_5of5_1790084828778.png")
    if os.path.exists(img_v5):
        add_card(s8, Inches(9.8), Inches(2.0), Inches(2.7), Inches(4.8), fill_color=COLOR_CARD, border_color=COLOR_GREEN)
        s8.shapes.add_picture(img_v5, Inches(9.85), Inches(2.05), Inches(2.6), Inches(4.7))

    add_speaker_notes(s8, """Judges, early-stage AI projects often look compelling in slides but fail when tested.

VERIFYX is grounded in verified software implementation:
- 24 out of 24 automated unit tests pass, validating edge cases, minimum object sizes (<32px), temporal confirmation, and report generation.
- Production build passes cleanly with zero TypeScript errors.
- Responsive layouts are validated across mobile and desktop browser viewports.
- Code is committed in Git under hash afb0952 and actively deployed on Vercel.

You can visit verifyx-gray.vercel.app on your own device and inspect the live product.""")

    # ====================================================
    # SLIDE 9 — IMPACT & FUTURE
    # ====================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, prs)
    add_header(s9, "STAGE 09 // ROADMAP & FUTURE WORKFLOWS", "From verification prototype to inspection platform.", "Clear separation between current live capabilities and future development horizons")

    h_w = Inches(3.7)
    h_h = Inches(3.8)
    h_top = Inches(2.0)

    c_h1 = add_card(s9, Inches(0.8), h_top, h_w, h_h, fill_color=COLOR_CARD, border_color=COLOR_GREEN)
    tf_h1 = c_h1.text_frame
    tf_h1.word_wrap = True
    p1 = tf_h1.paragraphs[0]
    p1.text = "TODAY // LIVE\nIMPLEMENTED & TESTED"
    p1.font.name = FONT_MONO
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_GREEN

    p_body = tf_h1.add_paragraph()
    p_body.space_before = Pt(14)
    p_body.text = "• In-browser verification web application\n• YOLO11n ONNX via WebAssembly\n• Spatial Rule Engine (clearance & size)\n• Tiered fallback pipeline (TF.js / Sim)\n• Closed-loop Before/After Rescan\n• Structured JSON inspection reports\n• Deployed on Vercel Edge Network"
    p_body.font.name = FONT_BODY
    p_body.font.size = Pt(12)
    p_body.font.color.rgb = COLOR_WHITE

    c_h2 = add_card(s9, Inches(4.8), h_top, h_w, h_h, fill_color=COLOR_CARD, border_color=COLOR_CYAN)
    tf_h2 = c_h2.text_frame
    tf_h2.word_wrap = True
    p1 = tf_h2.paragraphs[0]
    p1.text = "NEXT // ROADMAP\nPLANNED (NOT YET IMPLEMENTED)"
    p1.font.name = FONT_MONO
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_CYAN

    p_body = tf_h2.add_paragraph()
    p_body.space_before = Pt(14)
    p_body.text = "• REST & Webhook APIs for facility systems\n• Additional inspection categories\n• User-configurable spatial rule sets\n• Browser extension for desktop verification\n• Native iOS / Android wrapper packaging\n• Export integrations for safety databases"
    p_body.font.name = FONT_BODY
    p_body.font.size = Pt(12)
    p_body.font.color.rgb = COLOR_MUTED

    c_h3 = add_card(s9, Inches(8.8), h_top, h_w, h_h, fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER)
    tf_h3 = c_h3.text_frame
    tf_h3.word_wrap = True
    p1 = tf_h3.paragraphs[0]
    p1.text = "FUTURE // VISION\nRESEARCH (LONG-TERM)"
    p1.font.name = FONT_MONO
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_MUTED

    p_body = tf_h3.add_paragraph()
    p_body.space_before = Pt(14)
    p_body.text = "• Cryptographic signing for tamper-evident records\n• Multi-camera temporal tracking across zones\n• Automated compliance export adapters\n• Continuous on-device model adaptation\n• Universal verification protocol"
    p_body.font.name = FONT_BODY
    p_body.font.size = Pt(12)
    p_body.font.color.rgb = COLOR_MUTED

    disc_box = add_card(s9, Inches(0.8), Inches(6.05), Inches(11.7), Inches(0.8), fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER)
    tf_disc = disc_box.text_frame
    p = tf_disc.paragraphs[0]
    p.text = "TARGET USE CASES: Egress corridor clearance, emergency equipment accessibility, facility maintenance records."
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p2 = tf_disc.add_paragraph()
    p2.space_before = Pt(2)
    p2.text = "DISCLAIMER: VERIFYX is currently a prototype verification rule engine and not a certified safety compliance system."
    p2.font.name = FONT_BODY
    p2.font.size = Pt(10)
    p2.font.color.rgb = COLOR_MUTED

    add_speaker_notes(s9, """Where can VERIFYX go from here?

Today, we have demonstrated a working core thesis: explainable, in-browser visual verification with closed-loop rescan tracking.

Looking ahead, we plan to develop external webhook integrations, support custom user-defined spatial rules, and explore cryptographic signing for tamper-evident audit records.

Our goal is to build an accessible, explainable verification platform for routine inspection workflows.""")

    # ====================================================
    # SLIDE 10 — CLOSING
    # ====================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, prs)

    tag10 = s10.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.35))
    p = tag10.text_frame.paragraphs[0]
    p.text = "STAGE 10 // CALL TO ACTION · LIVE EVALUATION"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    c_stmt = s10.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(7.5), Inches(2.2))
    tf_st = c_stmt.text_frame
    tf_st.word_wrap = True
    p1 = tf_st.paragraphs[0]
    p1.text = "DON'T JUST CONSUME INFORMATION."
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p2 = tf_st.add_paragraph()
    p2.text = "VERIFY IT."
    p2.font.name = FONT_HEADING
    p2.font.size = Pt(44)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_CYAN

    lockup = s10.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(7.5), Inches(1.2))
    tf_lk = lockup.text_frame
    p_vx = tf_lk.paragraphs[0]
    p_vx.text = "VERIFYX"
    p_vx.font.name = FONT_HEADING
    p_vx.font.size = Pt(32)
    p_vx.font.bold = True
    p_vx.font.color.rgb = COLOR_WHITE
    p_sub = tf_lk.add_paragraph()
    p_sub.text = '"Verify Before You Trust."'
    p_sub.font.name = FONT_HEADING
    p_sub.font.size = Pt(20)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_MUTED

    pills_box = s10.shapes.add_textbox(Inches(0.8), Inches(5.1), Inches(7.5), Inches(0.8))
    tf_pil = pills_box.text_frame
    p = tf_pil.paragraphs[0]
    p.text = "● LIVE WEB APPLICATION     ● 24/24 TESTS PASSED     ● DEVICE-LOCAL PROCESSING"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN

    qr_card = add_card(s10, Inches(8.8), Inches(1.3), Inches(3.7), Inches(5.0), fill_color=COLOR_CARD, border_color=COLOR_CYAN)
    tf_qr = qr_card.text_frame
    tf_qr.word_wrap = True
    p = tf_qr.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "SCAN FOR LIVE DEMO"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    if os.path.exists(qr_path):
        s10.shapes.add_picture(qr_path, Inches(9.25), Inches(1.95), Inches(2.8), Inches(2.8))

    qr_lbl = s10.shapes.add_textbox(Inches(8.8), Inches(4.9), Inches(3.7), Inches(1.2))
    tf_lbl = qr_lbl.text_frame
    tf_lbl.word_wrap = True
    p1 = tf_lbl.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    p1.text = "https://verifyx-gray.vercel.app/"
    p1.font.name = FONT_MONO
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p2 = tf_lbl.add_paragraph()
    p2.space_before = Pt(4)
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "Runs directly on your smartphone browser.\nNo app download required."
    p2.font.name = FONT_BODY
    p2.font.size = Pt(9)
    p2.font.color.rgb = COLOR_MUTED

    f10 = s10.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.4))
    p = f10.text_frame.paragraphs[0]
    p.text = "BUILT • TESTED • DEPLOYED    |    GITHUB COMMIT: afb0952"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_MUTED

    add_speaker_notes(s10, """In an environment where visual assumptions lead to overlooked hazards, having an explainable verification workflow matters.

Don't assume.
Verify before you trust.

The product is live right now at verifyx-gray.vercel.app. Visit the link or scan the QR code to test VERIFYX directly on your phone.

Thank you, and I look forward to your questions.""")

    output_pptx = "c:/Users/heman/OneDrive/Desktop/VERIFYX/verifyx_pitch_deck.pptx"
    prs.save(output_pptx)
    print(f"Verified presentation saved successfully to: {output_pptx}")

if __name__ == "__main__":
    build_deck()
